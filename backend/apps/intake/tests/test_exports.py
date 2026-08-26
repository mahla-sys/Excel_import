"""Tests for secure XLSX/JSON exports (phase 6)."""
import io
import json

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.intake.exports import (
    CANONICAL_GUARDIAN_HEADERS,
    CANONICAL_STUDENT_HEADERS,
    ERROR_HEADERS,
    build_canonical_json,
    build_canonical_xlsx,
    build_error_report_xlsx,
    safe_cell,
)
from apps.intake.normalizers import CANONICAL_SCHEMA_VERSION

pytestmark = pytest.mark.django_db


def _xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "ثبت‌نام"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _process_via_api(api_client, operator):
    rows = [
        ["نام", "نام خانوادگی", "کدملی", "پایه"],
        ["علی", "رضایی", "0012345679", "هفتم"],
        ["ناقص", "بدون‌کدملی", "", "هفتم"],
    ]
    f = SimpleUploadedFile(
        "s.xlsx", _xlsx(rows),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    api_client.force_authenticate(operator)
    bid = api_client.post(
        "/api/v1/intake/batches/upload/",
        data={"file": f, "academic_year_label": "1405-1406"},
        format="multipart", SERVER_NAME="localhost",
    ).data["id"]
    api_client.post(
        f"/api/v1/intake/batches/{bid}/select-header/",
        data={"header_row": 1}, format="json", SERVER_NAME="localhost",
    )
    mappings = {
        "student.first_name": {"source_column_index": 0, "source_column_label": "نام"},
        "student.last_name": {"source_column_index": 1, "source_column_label": "نام خانوادگی"},
        "student.national_id": {"source_column_index": 2, "source_column_label": "کدملی"},
        "enrollment.grade": {"source_column_index": 3, "source_column_label": "پایه"},
    }
    api_client.post(
        f"/api/v1/intake/batches/{bid}/mapping/",
        data={"mappings": mappings}, format="json", SERVER_NAME="localhost",
    )
    api_client.post(
        f"/api/v1/intake/batches/{bid}/process/",
        format="json", SERVER_NAME="localhost",
    )
    return bid


def test_safe_cell_blocks_formula_triggers():
    assert safe_cell("=1+1").startswith("'")
    assert safe_cell("+CMD").startswith("'")
    assert safe_cell("-2+2").startswith("'")
    assert safe_cell("@SUM(A1)").startswith("'")
    # Non-formula strings pass through unchanged.
    assert safe_cell("علی") == "علی"
    # Numbers/dates are not prefixed.
    assert safe_cell(5) == 5


def test_error_report_only_includes_issue_rows(api_client, operator_a):
    bid = _process_via_api(api_client, operator_a)
    batch = __import__(
        "apps.intake.models", fromlist=["ImportBatch"]
    ).ImportBatch.objects.get(pk=bid)

    content = build_error_report_xlsx(batch)
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    # Header + both students (one has NID/incomplete errors, the other has
    # a missing-emergency-contact warning).
    assert rows[0] == tuple(ERROR_HEADERS)
    assert len(rows) == 3


def test_canonical_xlsx_has_two_sheets(api_client, operator_a):
    bid = _process_via_api(api_client, operator_a)
    from apps.intake.models import ImportBatch
    batch = ImportBatch.objects.get(pk=bid)
    content = build_canonical_xlsx(batch)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["دانش‌آموزان", "والدین"]

    student_rows = list(wb["دانش‌آموزان"].iter_rows(values_only=True))
    assert student_rows[0] == tuple(CANONICAL_STUDENT_HEADERS)
    assert len(student_rows) == 3  # header + 2 students

    guardian_rows = list(wb["والدین"].iter_rows(values_only=True))
    assert guardian_rows[0] == tuple(CANONICAL_GUARDIAN_HEADERS)


def test_canonical_json_is_versioned_and_school_scoped(api_client, operator_a):
    bid = _process_via_api(api_client, operator_a)
    from apps.intake.models import ImportBatch
    batch = ImportBatch.objects.get(pk=bid)
    payload = build_canonical_json(batch)

    assert payload["schema_version"] == CANONICAL_SCHEMA_VERSION
    assert payload["school"]["external_code"] == operator_a.school_tenant.external_code
    # Internal portal IDs of other systems must NOT be used as cross-system IDs.
    assert "school_id" not in payload["school"]
    assert payload["batch"]["file_sha256"] == batch.file_hash
    assert len(payload["students"]) == 2

    s0 = payload["students"][0]
    assert s0["identity"]["national_id"] == "0012345679"
    assert s0["enrollment"]["academic_year"] == "1405-1406"
    assert "raw_data" not in s0  # canonical export contains cleaned data


def test_export_endpoints_are_tenant_scoped(api_client, operator_a, operator_b):
    bid = _process_via_api(api_client, operator_a)
    api_client.force_authenticate(operator_b)
    for url in [
        f"/api/v1/intake/batches/{bid}/export/errors/",
        f"/api/v1/intake/batches/{bid}/export/canonical/xlsx/",
        f"/api/v1/intake/batches/{bid}/export/canonical/json/",
    ]:
        resp = api_client.get(url, SERVER_NAME="localhost")
        assert resp.status_code == 404, f"{url} should be 404 for other school"


def test_export_writes_audit_log(api_client, operator_a):
    bid = _process_via_api(api_client, operator_a)
    api_client.get(
        f"/api/v1/intake/batches/{bid}/export/canonical/json/",
        SERVER_NAME="localhost",
    )
    assert AuditLog.objects.filter(
        entity_type="ImportBatch",
        entity_id=str(bid),
        action="export.canonical",
    ).exists()


def test_export_blocks_until_processed(api_client, operator_a):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "s"
    ws.append(["نام"])
    buf = io.BytesIO()
    wb.save(buf)
    f = SimpleUploadedFile(
        "u.xlsx", buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    api_client.force_authenticate(operator_a)
    bid = api_client.post(
        "/api/v1/intake/batches/upload/",
        data={"file": f, "academic_year_label": "1405-1406"},
        format="multipart", SERVER_NAME="localhost",
    ).data["id"]
    resp = api_client.get(
        f"/api/v1/intake/batches/{bid}/export/canonical/json/",
        SERVER_NAME="localhost",
    )
    assert resp.status_code == 400
