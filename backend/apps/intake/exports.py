"""
Exports for an ImportBatch.

Three formats are produced:

1. **Error report (.xlsx)** — one row per StagedStudent that has at least
   one issue, designed for the school operator to review.
2. **Canonical workbook (.xlsx)** — one row per StagedStudent with all
   standardized identity/enrollment/contact columns. Guardians are on a
   second sheet.
3. **Canonical JSON** — versioned payload (``schema_version``) that is the
   stable contract the future integration connector reads.

Security:
* Every text cell that begins with a spreadsheet formula trigger
  (``=``, ``+``, ``-``, ``@``, tab, CR) is prefixed with an apostrophe so
  Excel does not execute it (CSV/Formula Injection / CWE-1236).
* The endpoints inherit the school-scoped queryset so an operator can only
  export data from their own school.
* Every export is recorded in the AuditLog.
"""
from __future__ import annotations

import json
from datetime import date, datetime
from typing import Any, Iterable
from urllib.parse import quote

from copy import copy

from django.http import HttpResponse
from openpyxl import Workbook

from .constants import RowStatus
from .models import ImportBatch, StagedStudent
from .normalizers import CANONICAL_SCHEMA_VERSION


# Characters that could cause spreadsheet applications to treat a cell as
# a formula. See CWE-1236.
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r", "\n")


def safe_cell(value: Any) -> Any:
    """
    Make a value safe to put in a spreadsheet cell.

    Non-string values (numbers, dates) are returned as-is. Strings that
    start with a formula trigger are prefixed with a single quote so
    Excel treats them as literal text.
    """
    if value is None:
        return ""
    if isinstance(value, (int, float, date, datetime)):
        return value
    text = str(value)
    if text and text[0] in _FORMULA_TRIGGERS:
        return "'" + text
    return text


def _status_label(status: str) -> str:
    return dict(RowStatus.choices).get(status, status)


# --------------------------------------------------------------------------
# Error report
# --------------------------------------------------------------------------

ERROR_HEADERS = [
    "ردیف",
    "نام",
    "نام خانوادگی",
    "کد ملی",
    "پایه",
    "وضعیت",
    "خطاها",
    "هشدارها",
    "پیشنهاد اقدام",
]


def build_error_report_xlsx(batch: ImportBatch) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش خطاها"

    ws.append(ERROR_HEADERS)
    _bold_header(ws)

    for s in batch.staged_students.all().prefetch_related("guardians"):
        issues = s.validation_issues or []
        errors = [i for i in issues if i.get("severity") == "error"]
        warnings = [i for i in issues if i.get("severity") == "warning"]
        if not issues:
            continue
        suggestion = _suggest_action(issues, s)
        row_number = s.source_row.row_number if s.source_row_id else ""
        ws.append([
            safe_cell(row_number),
            safe_cell(s.first_name),
            safe_cell(s.last_name),
            safe_cell(s.national_id),
            safe_cell(s.grade),
            safe_cell(_status_label(s.status)),
            safe_cell("\n".join(f"• {e.get('message') or e.get('code')}" for e in errors)),
            safe_cell("\n".join(f"• {w.get('message') or w.get('code')}" for w in warnings)),
            safe_cell(suggestion),
        ])

    _autosize(ws)
    return _save_virtual(wb)


def _suggest_action(issues: list[dict], student: StagedStudent) -> str:
    codes = {i.get("code") for i in issues}
    actions: list[str] = []
    if "W-NID-MISSING" in codes:
        actions.append("کد ملی را از پرونده کاغذی تکمیل کنید.")
    if any(c.startswith("E-NID") for c in codes):
        actions.append("صحت کد ملی را با شناسنامه تطبیق دهید.")
    if "W-DUPLICATE-EXACT" in codes or "W-DUPLICATE-SUSPECTED" in codes:
        actions.append("در صفحه‌ی «موارد نیازمند بررسی» رفع ابهام کنید.")
    if "W-EMERGENCY-CONTACT-MISSING" in codes:
        actions.append("شماره تماس ضروری ثبت شود.")
    if "E-BIRTHDATE-INVALID" in codes:
        actions.append("تاریخ تولد را به فرمت 1380/06/15 تصحیح کنید.")
    if "W-GRADE-MISSING" in codes:
        actions.append("پایه مشخص شود.")
    if not actions:
        actions.append("در صفحه‌ی بررسی پرونده، فیلد مربوطه را اصلاح کنید.")
    return " ".join(actions)


# --------------------------------------------------------------------------
# Canonical XLSX
# --------------------------------------------------------------------------

CANONICAL_STUDENT_HEADERS = [
    "ردیف",
    "نام",
    "نام خانوادگی",
    "نام کامل",
    "کد ملی",
    "کد ملی معتبر",
    "کد پرونده",
    "تاریخ تولد",
    "جنسیت",
    "محل تولد",
    "سال تحصیلی",
    "پایه",
    "کلاس/شعبه",
    "وضعیت کلاس‌بندی",
    "وضعیت ثبت‌نام",
    "موبایل دانش‌آموز",
    "تماس ضروری",
    "نسبت تماس ضروری",
    "تلفن ثابت",
    "سهمیه",
    "مدرسه قبلی",
    "وضعیت",
]

CANONICAL_GUARDIAN_HEADERS = [
    "دانش‌آموز",
    "نسبت",
    "نام",
    "نام خانوادگی",
    "نام کامل",
    "کد ملی",
    "تاریخ تولد",
    "موبایل",
    "تحصیلات",
    "شغل",
]


def build_canonical_xlsx(batch: ImportBatch) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "دانش‌آموزان"
    ws.append(CANONICAL_STUDENT_HEADERS)
    _bold_header(ws)

    contacts_default: dict = {}

    for s in batch.staged_students.all().select_related("source_row"):
        contacts = {**contacts_default, **(s.contacts or {})}
        ws.append([
            safe_cell(s.source_row.row_number if s.source_row_id else ""),
            safe_cell(s.first_name),
            safe_cell(s.last_name),
            safe_cell(s.full_name),
            safe_cell(s.national_id),
            safe_cell("بله" if s.national_id_valid else "خیر" if s.national_id_valid is False else ""),
            safe_cell(s.source_student_code),
            safe_cell(s.birth_date.isoformat() if s.birth_date else ""),
            safe_cell(s.gender),
            safe_cell(s.birth_place),
            safe_cell(s.academic_year),
            safe_cell(s.grade),
            safe_cell(s.classroom_title),
            safe_cell(s.get_placement_status_display() if s.placement_status else ""),
            safe_cell(s.get_admission_status_display() if s.admission_status else ""),
            safe_cell(contacts.get("student_mobile", "")),
            safe_cell(contacts.get("emergency_mobile", "")),
            safe_cell(contacts.get("emergency_contact_relation", "")),
            safe_cell(contacts.get("home_phone", "")),
            safe_cell((s.admission_data or {}).get("quota", "")),
            safe_cell((s.admission_data or {}).get("previous_school_name", "")),
            safe_cell(_status_label(s.status)),
        ])

    ws2 = wb.create_sheet("والدین")
    ws2.append(CANONICAL_GUARDIAN_HEADERS)
    for cell in ws2[1]:
        cell.font = copy(cell.font); cell.font = cell.font.copy(bold=True)

    for s in batch.staged_students.all().prefetch_related("guardians"):
        for g in s.guardians.all():
            ws2.append([
                safe_cell(s.full_name),
                safe_cell(g.get_relation_type_display()),
                safe_cell(g.first_name),
                safe_cell(g.last_name),
                safe_cell(g.full_name),
                safe_cell(g.national_id),
                safe_cell(g.birth_date.isoformat() if g.birth_date else ""),
                safe_cell(g.mobile),
                safe_cell(g.education_level),
                safe_cell(g.occupation),
            ])

    _autosize(ws)
    _autosize(ws2)
    return _save_virtual(wb)


# --------------------------------------------------------------------------
# Canonical JSON
# --------------------------------------------------------------------------

def build_canonical_json(batch: ImportBatch) -> dict:
    school = batch.school
    students_out = []
    for s in batch.staged_students.all().select_related("source_row").prefetch_related("guardians"):
        contacts = s.contacts or {}
        admission = s.admission_data or {}
        students_out.append({
            "id": s.id,
            "identity": {
                "first_name": s.first_name,
                "last_name": s.last_name,
                "full_name": s.full_name,
                "national_id": s.national_id,
                "national_id_valid": s.national_id_valid,
                "source_student_code": s.source_student_code,
                "birth_date": s.birth_date.isoformat() if s.birth_date else None,
                "birth_date_raw": s.birth_date_raw,
                "gender": s.gender,
                "birth_place": s.birth_place,
            },
            "enrollment": {
                "academic_year": s.academic_year,
                "grade": s.grade,
                "classroom_title": s.classroom_title,
                "placement_status": s.placement_status,
                "admission_status": s.admission_status,
            },
            "contacts": {
                "student_mobile": contacts.get("student_mobile", ""),
                "emergency_mobile": contacts.get("emergency_mobile", ""),
                "emergency_contact_relation": contacts.get("emergency_contact_relation", ""),
                "home_phone": contacts.get("home_phone", ""),
                "home_address": contacts.get("home_address", ""),
                "postal_code": contacts.get("postal_code", ""),
            },
            "guardians": [
                {
                    "relation_type": g.relation_type,
                    "first_name": g.first_name,
                    "last_name": g.last_name,
                    "full_name": g.full_name,
                    "national_id": g.national_id,
                    "birth_date": g.birth_date.isoformat() if g.birth_date else None,
                    "mobile": g.mobile,
                    "occupation": g.occupation,
                    "education_level": g.education_level,
                }
                for g in s.guardians.all()
            ],
            "admission": {
                "quota": admission.get("quota", ""),
                "previous_school_name": admission.get("previous_school_name", ""),
            },
            "extra_attributes": s.extra_attributes or {},
            "validation_issues": s.validation_issues or [],
            "status": s.status,
            "source_row_number": s.source_row.row_number if s.source_row_id else None,
        })

    return {
        "schema_version": CANONICAL_SCHEMA_VERSION,
        "exported_at": datetime.utcnow().isoformat() + "Z",
        "school": {
            # The external_code is the stable cross-system identifier. The
            # portal's internal PK is NOT a contract with the main system.
            "external_code": school.external_code,
            "name": school.name,
        },
        "batch": {
            "id": batch.id,
            "academic_year_label": batch.academic_year_label,
            "original_filename": batch.original_filename,
            "file_sha256": batch.file_hash,
            "uploaded_at": batch.uploaded_at.isoformat() if batch.uploaded_at else None,
            "processed_at": batch.processed_at.isoformat() if batch.processed_at else None,
        },
        "students": students_out,
    }


# --------------------------------------------------------------------------
# HTTP response helpers
# --------------------------------------------------------------------------

def xlsx_response(content: bytes, filename: str) -> HttpResponse:
    response = HttpResponse(
        content,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        f"attachment; filename=\"{filename}\"; "
        f"filename*=UTF-8''{quote(filename)}"
    )
    response["Content-Length"] = str(len(content))
    response["X-Content-Type-Options"] = "nosniff"
    return response


def json_response(payload: dict, filename: str) -> HttpResponse:
    content = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    response = HttpResponse(content, content_type="application/json; charset=utf-8")
    response["Content-Disposition"] = (
        f"attachment; filename=\"{filename}\"; "
        f"filename*=UTF-8''{quote(filename)}"
    )
    response["X-Content-Type-Options"] = "nosniff"
    return response


# --------------------------------------------------------------------------
# Internals
# --------------------------------------------------------------------------

def _bold_header(ws) -> None:
    for cell in ws[1]:
        cell.font = copy(cell.font)
        cell.font = cell.font.copy(bold=True)


def _save_virtual(wb: Workbook) -> bytes:
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _autosize(ws) -> None:
    """Set column widths based on the longest cell value (bounded)."""
    for col in ws.columns:
        max_len = 8
        letter = col[0].column_letter
        for cell in col:
            v = cell.value
            if v is None:
                continue
            length = max(len(line) for line in str(v).splitlines()) if str(v) else 0
            if length > max_len:
                max_len = length
        ws.column_dimensions[letter].width = min(max_len + 2, 50)
